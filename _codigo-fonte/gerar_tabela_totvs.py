"""Leitura do DadosEntregues / tabelas TOTVS antigas e gravacao da Tabela Controle Correios Itu - TOTVS.

Usado pelo programa (servidor.py). Pastas:
  PASTA     - onde fica o .exe (ou a pasta acima de _codigo-fonte): Modelos, configuracoes
  RECURSOS  - arquivos que vao dentro do .exe: interface.html e o Modelo
  SAIDA     - pasta dos dados (historico e tabelas); pode ser trocada para o servidor
"""
import copy
import csv
import io
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

import openpyxl

if getattr(sys, "frozen", False):
    # rodando como .exe: dados ficam ao lado do .exe, arquivos internos no pacote
    PASTA = Path(sys.executable).resolve().parent
    RECURSOS = Path(sys._MEIPASS)
else:
    # rodando pelo codigo-fonte: os .py ficam em _codigo-fonte, um nivel abaixo do .exe
    RECURSOS = Path(__file__).resolve().parent
    PASTA = RECURSOS.parent
MODELO = PASTA / "Modelos" / "Modelo de Tabela TOTVS.xlsx"
if not MODELO.exists():
    MODELO = RECURSOS / "Modelo de Tabela TOTVS.xlsx"
SAIDA = PASTA / "Saida"

MESES = ["Janeiro", "Fevereiro", "Marco", "Abril", "Maio", "Junho", "Julho",
         "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]
VENDEDOR_ITU = re.compile(r"-\s*ITU\s*$", re.IGNORECASE)
RASTREIO_OK = re.compile(r"^[A-Z]{2}\d{9}[A-Z]{2}$")

# grafias encontradas nas planilhas antigas -> status padrao
STATUS_PADRAO = {
    "entregue": "Entregue",
    "entrgeue": "Entregue",
    "pedido ja entregue": "Entregue",
    "retirada na agencia": "Retirado na Agência",
    "retirado na agencia": "Retirado na Agência",
    "retornou ao remetente": "Retornou ao Remetente",
    "retornou ao remtente": "Retornou ao Remetente",
    "entregue ao remetente": "Retornou ao Remetente",
    "cancelado": "Pedido Cancelado",
    "pedido cancelado": "Pedido Cancelado",
    "etiqueta expirada": "Etiqueta Expirada",
    "etiqueta emitida - expirada": "Etiqueta Expirada",
    "etiqueta emitida": "Etiqueta Emitida",
    "em rota": "Em trânsito",
    "trocou de codigo": "Trocou de Código",
}


def _norm(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    return " ".join("".join(c for c in s if not unicodedata.combining(c)).lower().split())


def normalizar_status(s):
    s = " ".join(str(s or "").split())
    return STATUS_PADRAO.get(_norm(s), s)


def normalizar_vendedor(s):
    return " ".join(str(s or "").split()).upper()


def _linhas_do_arquivo(nome, conteudo):
    if nome.lower().endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    for enc in ("utf-8-sig", "cp1252"):
        try:
            texto = conteudo.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    return list(csv.reader(io.StringIO(texto), delimiter=";"))


def _data(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", str(v or ""))
    if not m:
        return None
    try:
        return date(int(m[3]), int(m[2]), int(m[1]))
    except ValueError:
        return None


def _erp(v):
    return str(int(v)) if isinstance(v, float) else str(v or "").strip()


def _achar_cabecalho(linhas):
    inicio = next((i for i, r in enumerate(linhas) if any(_norm(c) == "erp" for c in r)), None)
    if inicio is None:
        raise ValueError("Nao achei o cabecalho (coluna ERP). Confira se e o relatorio de dados entregues ou uma tabela TOTVS.")
    return inicio, [_norm(c) for c in linhas[inicio]]


def ler_arquivo(nome, conteudo):
    """Le um DadosEntregues ou uma Tabela TOTVS ja preenchida.

    Devolve (tipo, envios_validos, sem_rastreio); tipo e "dados" ou "tabela".
    """
    linhas = _linhas_do_arquivo(nome, conteudo)
    _, cab = _achar_cabecalho(linhas)
    if "nome vended." in cab:
        return ("dados",) + _ler_dados_entregues(linhas)
    return ("tabela",) + _ler_tabela_totvs(nome, linhas)


def _ler_dados_entregues(linhas):
    inicio, cab = _achar_cabecalho(linhas)
    idx = {}
    for chave, coluna in (("erp", "erp"), ("vendedor", "nome vended."), ("cliente", "cliente"),
                          ("data", "dt.emissao"), ("rastreio", "codigo rastreio")):
        if coluna not in cab:
            raise ValueError(f'Coluna "{coluna}" nao encontrada no arquivo.')
        idx[chave] = cab.index(coluna)

    validos, sem_rastreio = [], []
    for r in linhas[inicio + 1:]:
        r = list(r) + [None] * (len(cab) - len(r))
        vendedor = normalizar_vendedor(r[idx["vendedor"]])
        if not VENDEDOR_ITU.search(vendedor):
            continue
        dia = _data(r[idx["data"]])
        if not dia:
            continue
        envio = {
            "erp": _erp(r[idx["erp"]]),
            "cliente": str(r[idx["cliente"]] or "").strip(),
            "vendedor": vendedor,
            "rastreio": str(r[idx["rastreio"]] or "").strip().upper(),
            "data": dia.isoformat(),
        }
        (validos if RASTREIO_OK.match(envio["rastreio"]) else sem_rastreio).append(envio)
    return validos, sem_rastreio


def _ler_tabela_totvs(nome, linhas):
    inicio, cab = _achar_cabecalho(linhas)

    def col(*nomes):
        return next((cab.index(n) for n in nomes if n in cab), None)

    idx = {
        "data": col("data"), "erp": col("erp"), "cliente": col("cliente", "cleinte"),
        "vendedor": col("vendedor", "vedndedor"), "rastreio": col("rastreio", "registro"),
        "status": col("status"),
    }
    if idx["rastreio"] is None:
        raise ValueError(f"{nome}: coluna RASTREIO/REGISTRO nao encontrada.")

    # dia da tabela: titulo "CORREIOS ITU - TOTVS - dd/mm/aaaa" na primeira linha
    dia_titulo = next((_data(c) for r in linhas[:inicio] for c in r if _data(c)), None)

    def pega(r, k):
        return r[idx[k]] if idx[k] is not None and idx[k] < len(r) else None

    validos, sem_rastreio = [], []
    for r in linhas[inicio + 1:]:
        rastreio = str(pega(r, "rastreio") or "").strip().upper().replace(" ", "")
        erp = _erp(pega(r, "erp"))
        if not rastreio and not erp:
            continue
        dia = _data(pega(r, "data")) or dia_titulo
        if not dia:
            continue
        envio = {
            "erp": erp,
            "cliente": " ".join(str(pega(r, "cliente") or "").split()),
            "vendedor": normalizar_vendedor(pega(r, "vendedor")),
            "rastreio": rastreio,
            "data": dia.isoformat(),
            "status": normalizar_status(pega(r, "status")),
        }
        (validos if RASTREIO_OK.match(rastreio) else sem_rastreio).append(envio)
    return validos, sem_rastreio


def caminho_tabela(dia, base=None):
    """<base>/2026/10 - Outubro/Tabela ... - 21 de Outubro.xlsx"""
    pasta = (base or SAIDA) / str(dia.year) / f"{dia.month:02d} - {MESES[dia.month - 1]}"
    return pasta / f"Tabela Controle Correios Itu - TOTVS - {dia.day:02d} de {MESES[dia.month - 1]}.xlsx"


def salvar_tabela(dia, envios, destino=None, titulo=None):
    """Grava a tabela TOTVS usando o Modelo como base.

    `dia` vale para o titulo e a data padrao; num relatorio de varios dias
    passe `titulo` e cada linha usa a data do proprio envio.
    """
    envios = sorted(envios, key=lambda e: (e["data"], int(e["erp"]) if e["erp"].isdigit() else 0))
    wb = openpyxl.load_workbook(MODELO)
    ws = wb.active
    ws["A1"] = titulo or f"CORREIOS ITU - TOTVS - {dia:%d/%m/%Y}"

    # cabecalho do modelo: DATA, ERP, CLIENTE, VENDEDOR, RASTREIO, STATUS
    estilo = [copy.copy(ws.cell(3, c)._style) for c in range(1, 7)]
    altura = ws.row_dimensions[3].height
    ws.delete_rows(3, ws.max_row)

    for i, e in enumerate(envios, start=3):
        d = date.fromisoformat(e["data"]) if titulo else dia
        valores = [
            datetime(d.year, d.month, d.day),
            int(e["erp"]) if e["erp"].isdigit() else e["erp"],
            e["cliente"] or None,  # sem cliente fica em branco
            e["vendedor"],
            e["rastreio"],
            e.get("status") or None,  # STATUS fica em branco ate ser preenchido
        ]
        for c, v in enumerate(valores, start=1):
            cel = ws.cell(i, c, v)
            cel._style = copy.copy(estilo[c - 1])
            cel.number_format = "dd/mm/yyyy" if c == 1 else "General"
        ws.row_dimensions[i].height = altura

    largura_cliente = max([10] + [len(e["cliente"] or "") + 2 for e in envios])
    for col, w in {"A": 12, "B": 9, "C": min(largura_cliente, 40), "D": 30, "E": 17, "F": 22}.items():
        ws.column_dimensions[col].width = w

    destino = destino or caminho_tabela(dia)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino
